"""중간 모델(md_parser) → docx. 「5. 서식 규칙」을 구현한다."""
from __future__ import annotations

import copy
import math
import os
import random
import re
from dataclasses import dataclass, field, replace
from urllib.parse import unquote

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.table import _Cell

import ooxml_helpers as ox
import styles as S
from md_parser import (CodeBlock, Heading, ListBlock, MathBlock, MdDocument, Paragraph,
                       Quote, Rule, Span, Table, spans_text)

# ── 정규식 ─────────────────────────────────────────────
EMOJI_RE = re.compile(
    "[\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\u2300-\u23FF\u2190-\u21FF\uFE0F\u200D"
    "\U000E0000-\U000E007F]+")
# ↔→ 같은 화살표는 맑은 고딕에 있으므로 이모지 글꼴로 바꾸지 않는다
_ARROWS = set("←↑→↓↔↕⇒⇔")
UNITS = ("개월|개국|퍼센트|시간|가지|단계|학점|원|초|분|일|주|년|명|개|건|회|배|점|위|살|세|"
         "쪽|곳|번|권|대|인|%")
NUM_RE = re.compile(
    rf"(?<![0-9A-Za-z.:/\-])(?:\d{{1,3}}(?:,\d{{3}})+|\d+(?:\.\d+)?)"
    rf"(?:\s*[~∼\-–]\s*\d+(?:\.\d+)?)?\s*(?:[만천억조]\s*)?(?:{UNITS})"
    rf"|(?<![A-Za-z])[A-D]\+(?![+A-Za-z])")
UNCERTAIN_RE = re.compile(r"\[\?\]")
EMDASH_RE = re.compile(r"(?<!\d)\s+—\s+(?!\d)")
PIPE_RE = re.compile(r"\s*\|\s*")
SENT_RE = re.compile(r"(?:(?<=[^\d.\s][.?!])|(?<=[.?!][)\"'”’\]]))\s+(?:/\s+)?")
FEED_RE = re.compile(r"(?:^|(?<=\s))▶\s*")
RESIDUE = ["|", "---", "**", "##", "`"]


# ── 글자 폭 추정 (줄 수·높이 계산용) ───────────────────
def text_em(s: str) -> float:
    w = 0.0
    for ch in s:
        o = ord(ch)
        if ch == " ":
            w += 0.3
        elif o < 0x80:
            w += 0.62 if ch.isupper() else 0.55
        elif (0xAC00 <= o <= 0xD7AF or 0x1100 <= o <= 0x11FF or 0x3130 <= o <= 0x318F
              or 0x4E00 <= o <= 0x9FFF or 0x3000 <= o <= 0x30FF or 0xFF00 <= o <= 0xFFEF):
            w += 0.95
        else:
            w += 1.0
    return w


def est_lines(text: str, width_tw: float, size: float) -> int:
    width_pt = max(width_tw / 20, 20)
    n = 0
    for line in text.split("\n"):
        n += max(1, math.ceil(text_em(line) * size / width_pt))
    return n


# ── 옵션과 결과 ────────────────────────────────────────
@dataclass
class BuildOptions:
    auto_number_bold: bool = True
    auto_box: bool = True
    update_fields_on_open: bool = True   # Word·LibreOffice 둘 다 없을 때 True
    mml2omml: str | None = None          # MML2OMML.XSL 경로
    for_pdf: bool = False                # PDF용 중간 문서면 이모지 글꼴을 PDF_EMOJI_FONT로


@dataclass
class HeadingInfo:
    bookmark: str
    text: str
    level: int          # 1=장, 2=소제목, 3=하위


@dataclass
class BuildResult:
    warnings: list[str] = field(default_factory=list)
    headings: list[HeadingInfo] = field(default_factory=list)
    title: str = ""


@dataclass
class Ctx:
    width: int                  # 사용 가능한 폭(twip)
    where: str = "body"         # body | box | cell
    base_left: int = 0          # 목록 안 블록의 왼쪽 들여쓰기
    pstyle: str = "Normal"

    @property
    def size(self):
        return S.TABLE_SIZE if self.where == "cell" else S.BODY_SIZE

    def sub(self, **kw):
        return replace(self, **kw)


def find_mml2omml() -> str | None:
    roots = [os.environ.get("ProgramFiles", r"C:\Program Files"),
             os.environ.get("ProgramFiles(x86)", r"C:\Program Files (x86)")]
    for root in roots:
        for ver in ("Office16", "Office15", "Office14"):
            for mid in (r"Microsoft Office\root", "Microsoft Office"):
                p = os.path.join(root, mid, ver, "MML2OMML.XSL")
                if os.path.exists(p):
                    return p
    return None


# ── 빌더 ───────────────────────────────────────────────
class DocxBuilder:
    def __init__(self, md: MdDocument, md_path: str, opts: BuildOptions | None = None):
        self.md = md
        self.md_path = md_path
        self.md_dir = os.path.dirname(os.path.abspath(md_path))
        self.opts = opts or BuildOptions()
        self.doc = docx.Document()
        self.part = self.doc.part
        self.ids = ox.IdCounter(1)
        self.sdt_ids = ox.IdCounter(random.randint(100000, 900000))
        self.result = BuildResult()
        self.used_cells: set = set()
        self.label_ps: set = set()              # 소제목성 레이블 단락(w:p)
        self.last_p = None
        self.bookmarks: dict[int, str] = {}     # id(Heading) → 책갈피 이름
        self.levels: dict[int, int] = {}        # id(Heading) → 문서 제목 수준(1~3)
        self.title_heading = None
        self._xslt = None
        self._xslt_failed = False
        self.emoji_font = S.PDF_EMOJI_FONT if self.opts.for_pdf else S.EMOJI_FONT

    # ━━ 진입점 ━━
    def build(self, out_path: str) -> BuildResult:
        self._setup_document()
        self._analyze_headings()
        title = self._doc_title()
        self.result.title = title
        blocks = self.md.blocks
        if self.opts.auto_box:
            blocks = self._auto_box(blocks)
        self._collect_toc(blocks)
        body = Ctx(width=S.BODY_W)

        self._cover(title)
        self._toc()
        self._page_start()
        self._render_blocks(self.doc, blocks, body)
        self._footnotes(body)
        self._appendix(body)
        self._header_footer(title)
        self._trim_trailing()
        self.result.warnings.extend(self._residue_check())
        self.doc.save(out_path)
        return self.result

    # ━━ 문서 설정 ━━
    def _setup_document(self):
        sec = self.doc.sections[0]
        sp = sec._sectPr
        ox.set_child(sp, ox.el("w:pgSz", w=S.PAGE_W, h=S.PAGE_H))
        ox.set_child(sp, ox.el("w:pgMar", top=S.MARGIN, right=S.MARGIN, bottom=S.MARGIN,
                               left=S.MARGIN, header=S.HEADER_DIST, footer=S.HEADER_DIST,
                               gutter=0))
        ox.set_child(sp, ox.el("w:docGrid", type="default", linePitch=360))
        self._setup_styles()
        self._setup_settings()
        self._setup_numbering()

    def _setup_styles(self):
        st = self.doc.styles.element
        # 모든 글꼴 지정에서 테마 글꼴을 없애고 맑은 고딕으로 고정 (eastAsia 포함)
        for rf in st.iter(qn("w:rFonts")):
            for a in list(rf.attrib):
                del rf.attrib[a]
            for a in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn("w:" + a), S.FONT)
        dd = st.find(qn("w:docDefaults"))
        rpd = dd.find(qn("w:rPrDefault"))
        if rpd is None:
            rpd = ox.el("w:rPrDefault")
            dd.insert(0, rpd)
        rpr = rpd.find(qn("w:rPr"))
        if rpr is None:
            rpr = ox.el("w:rPr")
            rpd.append(rpr)
        ox.set_child(rpr, ox.el("w:rFonts", ascii=S.FONT, hAnsi=S.FONT, eastAsia=S.FONT, cs=S.FONT))
        ox.set_child(rpr, ox.el("w:sz", val=round(S.BODY_SIZE * 2)))
        ox.set_child(rpr, ox.el("w:szCs", val=round(S.BODY_SIZE * 2)))
        ox.set_child(rpr, ox.el("w:lang", val="en-US", eastAsia="ko-KR", bidi="ar-SA"))
        ppd = dd.find(qn("w:pPrDefault"))
        if ppd is None:
            ppd = ox.el("w:pPrDefault")
            dd.append(ppd)
        ppr = ppd.find(qn("w:pPr"))
        if ppr is None:
            ppr = ox.el("w:pPr")
            ppd.append(ppr)
        ox.set_child(ppr, ox.el("w:spacing", before=0, after=0, line=round(S.BODY_SIZE * S.BODY_LINE * 20),
                                lineRule="exact"))

        self._pstyle("Normal", size=S.BODY_SIZE, line=S.BODY_SIZE * S.BODY_LINE,
                     after=S.PARA_AFTER, color=S.DARK_TEXT, widow=True, snap=False)
        for lvl, h in S.HEADINGS.items():
            sid = f"Heading{lvl}"
            self._pstyle(sid, size=h["size"], line=h["size"] * h["line"], before=h["before"],
                         after=h["after"], color=h["color"], bold=True, keep_next=True,
                         keep_lines=True, outline=lvl - 1, reset=True)
            s_el = self._style_el(sid)
            ppr = s_el.find(qn("w:pPr"))
            if lvl == 1:
                bdr = ox.el("w:pBdr")
                bdr.append(ox.el("w:bottom", val="single", sz=S.H1_RULE_SIZE, space=4,
                                 color=S.H1_RULE_COLOR))
                ox.set_child(ppr, bdr)
            if lvl == 3:
                ox.shade(ppr, S.H3_BG)
        self._pstyle("MDTable", size=S.TABLE_SIZE, line=S.TABLE_SIZE * S.TABLE_LINE,
                     after=S.CELL_PARA_AFTER, new=True, name="MD 표 본문")
        self._pstyle("MDBox", size=S.BODY_SIZE, line=S.BODY_SIZE * S.BODY_LINE,
                     after=S.BOX_PARA_AFTER, new=True, name="MD 상자 본문")
        self._pstyle("MDSpacer", size=2, line=S.SPACER_LINE, after=S.SPACER_AFTER, new=True,
                     name="MD 간격")
        self._pstyle("MDCode", size=S.CODE_SIZE, line=S.CODE_SIZE * 1.5, after=0, new=True,
                     name="MD 코드", font=S.CODE_FONT, shade=S.CODE_BG, keep_lines=True)
        self._pstyle("MDCoverTitle", size=S.COVER_TITLE_SIZE,
                     line=S.COVER_TITLE_SIZE * S.COVER_TITLE_LINE, before=S.COVER_TOP_SPACE,
                     after=12, color=S.NAVY, bold=True, new=True, name="MD 표지 제목")
        ppr = self._style_el("MDCoverTitle").find(qn("w:pPr"))
        bdr = ox.el("w:pBdr")
        bdr.append(ox.el("w:bottom", val="single", sz=S.COVER_RULE_SZ, space=8, color=S.NAVY))
        ox.set_child(ppr, bdr)
        self._pstyle("MDCoverSub", size=S.COVER_SUB_SIZE, line=S.COVER_SUB_SIZE * 1.5, after=6,
                     color=S.BLUE, new=True, name="MD 표지 부제")
        self._pstyle("MDTocTitle", size=S.TOC_TITLE_SIZE, line=S.TOC_TITLE_SIZE * 1.3, after=18,
                     color=S.NAVY, bold=True, new=True, name="MD 차례 제목")
        ppr = self._style_el("MDTocTitle").find(qn("w:pPr"))
        bdr = ox.el("w:pBdr")
        bdr.append(ox.el("w:bottom", val="single", sz=12, space=6, color=S.NAVY))
        ox.set_child(ppr, bdr)
        self._pstyle("MDToc1", size=S.TOC_CH_SIZE, line=S.TOC_CH_SIZE * S.BODY_LINE, after=0,
                     color=S.NAVY, bold=True, new=True, name="MD 차례 1")
        self._pstyle("MDToc2", size=S.TOC_SUB_SIZE, line=S.TOC_SUB_SIZE * S.BODY_LINE, after=0,
                     color=S.DARK_TEXT, new=True, name="MD 차례 2")
        self._pstyle("MDTocNote", size=S.TOC_NOTE_SIZE, line=S.TOC_NOTE_SIZE * S.BODY_LINE,
                     after=0, color=S.GRAY_TEXT, new=True, name="MD 차례 안내")
        self._pstyle("Header", size=S.HEADER_SIZE, line=S.HEADER_SIZE * 1.5, after=0,
                     color=S.GRAY_TEXT)
        self._pstyle("Footer", size=S.FOOTER_SIZE, line=S.FOOTER_SIZE * 1.5, after=0,
                     color=S.GRAY_TEXT)
        self._cstyle("MDCodeChar", "MD 인라인 코드", font=S.CODE_FONT, shade=S.CODE_BG)
        self._cstyle("MDLink", "MD 링크", color=S.LINK_COLOR, underline=True)

    def _style_el(self, sid):
        for s in self.doc.styles.element.findall(qn("w:style")):
            if s.get(qn("w:styleId")) == sid:
                return s
        return None

    def _pstyle(self, sid, size, line, before=0, after=0, color=None, bold=False, keep_next=False,
                keep_lines=False, outline=None, new=False, name=None, font=None, shade=None,
                widow=True, snap=None, reset=False):
        s = self._style_el(sid)
        if s is None:
            s = ox.el("w:style", type="paragraph", styleId=sid)
            s.append(ox.el("w:name", val=name or sid))
            s.append(ox.el("w:basedOn", val="Normal"))
            s.append(ox.el("w:qFormat"))
            self.doc.styles.element.append(s)
        if reset:
            for tag in ("w:pPr", "w:rPr"):
                e = s.find(qn(tag))
                if e is not None:
                    s.remove(e)
        ppr = s.find(qn("w:pPr"))
        if ppr is None:
            ppr = ox.el("w:pPr")
            rpr_existing = s.find(qn("w:rPr"))
            (rpr_existing.addprevious(ppr) if rpr_existing is not None else s.append(ppr))
        rpr = s.find(qn("w:rPr"))
        if rpr is None:
            rpr = ox.el("w:rPr")
            s.append(rpr)
        if keep_next:
            ox.set_child(ppr, ox.el("w:keepNext"))
        if keep_lines:
            ox.set_child(ppr, ox.el("w:keepLines"))
        if widow:
            ox.set_child(ppr, ox.el("w:widowControl"))
        if shade:
            ox.shade(ppr, shade)
        if snap is False:
            ox.set_child(ppr, ox.el("w:autoSpaceDE", val=0))
            ox.set_child(ppr, ox.el("w:autoSpaceDN", val=0))
            ox.set_child(ppr, ox.el("w:snapToGrid", val=0))
        ox.set_child(ppr, ox.el("w:spacing", before=ox.pt2tw(before), after=ox.pt2tw(after),
                                line=ox.pt2tw(line), lineRule="exact"))
        if outline is not None:
            ox.set_child(ppr, ox.el("w:outlineLvl", val=outline))
        f = font or S.FONT
        ox.set_child(rpr, ox.el("w:rFonts", ascii=f, hAnsi=f, eastAsia=S.FONT if font else f, cs=f))
        if bold:
            ox.set_child(rpr, ox.el("w:b"))
            ox.set_child(rpr, ox.el("w:bCs"))
        if color:
            ox.set_child(rpr, ox.el("w:color", val=color))
        ox.set_child(rpr, ox.el("w:sz", val=round(size * 2)))
        ox.set_child(rpr, ox.el("w:szCs", val=round(size * 2)))

    def _cstyle(self, sid, name, font=None, shade=None, color=None, underline=False):
        s = ox.el("w:style", type="character", styleId=sid)
        s.append(ox.el("w:name", val=name))
        s.append(ox.el("w:basedOn", val="DefaultParagraphFont"))
        rpr = ox.el("w:rPr")
        if font:
            ox.set_child(rpr, ox.el("w:rFonts", ascii=font, hAnsi=font, eastAsia=S.FONT, cs=font))
        if color:
            ox.set_child(rpr, ox.el("w:color", val=color))
        if underline:
            ox.set_child(rpr, ox.el("w:u", val="single"))
        if shade:
            ox.shade(rpr, shade)
        s.append(rpr)
        self.doc.styles.element.append(s)

    def _setup_settings(self):
        se = self.doc.settings.element
        compat = se.find(qn("w:compat"))
        if compat is not None:
            for cs in compat.findall(qn("w:compatSetting")):
                if cs.get(qn("w:name")) == "compatibilityMode":
                    cs.set(qn("w:val"), "15")
        tfl = se.find(qn("w:themeFontLang"))
        if tfl is not None:
            tfl.set(qn("w:eastAsia"), "ko-KR")
        if self.opts.update_fields_on_open:
            ox.set_child(se, ox.el("w:updateFields", val="true"))

    def _setup_numbering(self):
        numbering = self.part.numbering_part.element
        self.numbering = numbering
        abs_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
        num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        self._next_abs = max(abs_ids + [0]) + 1
        self._next_num = max(num_ids + [0]) + 1
        self.abs_bullet = self._add_abstract(bullet=True)
        self.abs_decimal = self._add_abstract(bullet=False)
        self.bullet_num = self._add_num(self.abs_bullet)

    def _add_abstract(self, bullet: bool) -> int:
        aid = self._next_abs
        self._next_abs += 1
        a = ox.el("w:abstractNum", abstractNumId=aid)
        a.append(ox.el("w:nsid", val="%08X" % random.getrandbits(32)))
        a.append(ox.el("w:multiLevelType", val="hybridMultilevel"))
        for lvl in range(9):
            lv = ox.el("w:lvl", ilvl=lvl)
            lv.append(ox.el("w:start", val=1))
            if bullet:
                lv.append(ox.el("w:numFmt", val="bullet"))
                lv.append(ox.el("w:lvlText", val=S.BULLETS[min(lvl, len(S.BULLETS) - 1)]))
            else:
                fmt, txt = [("decimal", "%{}."), ("decimal", "%{})"), ("lowerLetter", "%{}.")][lvl % 3]
                lv.append(ox.el("w:numFmt", val=fmt))
                lv.append(ox.el("w:lvlText", val=txt.format(lvl + 1)))
            lv.append(ox.el("w:lvlJc", val="left"))
            ppr = ox.el("w:pPr")
            ppr.append(ox.el("w:ind", left=S.LIST_INDENT * (lvl + 1), hanging=S.LIST_HANGING))
            lv.append(ppr)
            rpr = ox.el("w:rPr")
            rpr.append(ox.el("w:rFonts", ascii=S.FONT, hAnsi=S.FONT, eastAsia=S.FONT, cs=S.FONT,
                             hint="default"))
            lv.append(rpr)
            a.append(lv)
        first_num = self.numbering.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(a)
        else:
            self.numbering.append(a)
        return aid

    def _add_num(self, abs_id: int, restart_level: int | None = None, start: int = 1) -> int:
        nid = self._next_num
        self._next_num += 1
        n = ox.el("w:num", numId=nid)
        n.append(ox.el("w:abstractNumId", val=abs_id))
        if restart_level is not None:
            ov = ox.el("w:lvlOverride", ilvl=restart_level)
            ov.append(ox.el("w:startOverride", val=start))
            n.append(ov)
        self.numbering.append(n)
        return nid

    # ━━ 제목 분석 ━━
    def _all_headings(self, blocks):
        for b in blocks:
            if isinstance(b, Heading):
                yield b
            elif isinstance(b, Quote):
                yield from self._all_headings(b.blocks)
            elif isinstance(b, ListBlock):
                for it in b.items:
                    yield from self._all_headings(it.blocks)

    def _analyze_headings(self):
        hs = list(self._all_headings(self.md.blocks))
        h1 = [h for h in hs if h.level == 1]
        if len(h1) == 1 and hs and hs[0] is h1[0]:
            # 첫 # 하나뿐이면 문서 제목 → 본문에서는 빼고, ## 부터 제목1
            self.title_heading = h1[0]
            others = [h for h in hs if h is not self.title_heading]
            offset = (min(h.level for h in others) - 1) if others else 1
        else:
            offset = (min(h.level for h in hs) - 1) if hs else 0
        for h in hs:
            if h is self.title_heading:
                continue
            self.levels[id(h)] = max(1, min(3, h.level - offset))

    def _doc_title(self) -> str:
        t = self.md.meta.get("title")
        if t:
            return self._clean_text(str(t))
        if self.title_heading is not None:
            return self._clean_text(spans_text(self.title_heading.spans))
        return os.path.splitext(os.path.basename(self.md_path))[0]

    def _collect_toc(self, blocks):
        for h in self._all_headings(blocks):
            if h is self.title_heading:
                continue
            lvl = self.levels.get(id(h), 3)
            name = "_Toc%08d" % self.ids.n
            self.bookmarks[id(h)] = name
            self.ids.next()
            self.result.headings.append(
                HeadingInfo(name, self._clean_text(spans_text(h.spans).replace("\n", " ")), lvl))
        if self._appendix_items():
            name = "_Toc%08d" % self.ids.next()
            self._appendix_bm = name
            self.result.headings.append(HeadingInfo(name, S.APPENDIX_TITLE, 1))

    # ━━ 특정 절 자동 상자화 ━━
    def _auto_box(self, blocks):
        out, i = [], 0
        while i < len(blocks):
            b = blocks[i]
            kind = self._auto_box_kind(b) if isinstance(b, Heading) else None
            if not kind:
                out.append(b)
                i += 1
                continue
            out.append(b)
            i += 1
            lvl = b.level
            run = []

            def flush():
                trail = []
                while run and isinstance(run[-1], Rule):
                    trail.insert(0, run.pop())
                # 표는 상자에 중첩하지 않는다(큰 표가 든 단일 셀 행은 Word가 쪽에서
                # 나누지 못해 통째로 다음 쪽으로 밀린다). 대신 머리글 행을 상자 색으로 칠한다.
                chunk = []
                for blk in run + [None]:
                    if blk is None or isinstance(blk, Table):
                        if chunk:
                            if len(chunk) == 1 and isinstance(chunk[0], Quote):
                                out.append(chunk[0])
                            else:
                                out.append(self._make_box(kind, list(chunk)))
                            chunk = []
                        if blk is not None:
                            out.append(replace(blk, accent=kind))
                    else:
                        chunk.append(blk)
                out.extend(trail)
                run.clear()

            while i < len(blocks):
                nb = blocks[i]
                if isinstance(nb, Heading):
                    if nb.level <= lvl:
                        break
                    flush()
                    out.append(nb)
                else:
                    run.append(nb)
                i += 1
            flush()
        return out

    def _auto_box_kind(self, h: Heading):
        text = spans_text(h.spans)
        for words, kind in S.AUTO_BOX_SECTIONS:
            if all(w in text for w in words):
                return kind
        return None

    def _make_box(self, kind, blocks):
        title = None
        if blocks and isinstance(blocks[0], Paragraph) and self._is_label(blocks[0].spans):
            title = [replace(s, bold=False) for s in blocks[0].spans]
            blocks = blocks[1:]
        return Quote(blocks, kind, title)

    # ━━ 표지 ━━
    def _cover(self, title):
        p = self.doc.add_paragraph()
        ox.style(p._p, "MDCoverTitle")
        self._add_text(p._p, title, ctx=None, emoji=True)
        sub = self.md.meta.get("subtitle")
        if sub:
            p = self.doc.add_paragraph()
            ox.style(p._p, "MDCoverSub")
            self._add_text(p._p, self._clean_text(str(sub)), ctx=None, emoji=True)
        p = self.doc.add_paragraph()
        ox.style(p._p, "MDSpacer")
        ox.spacing(p._p, after=24)
        rows = []
        for key, label in S.COVER_KEYS:
            if key in self.md.meta and self.md.meta[key] not in (None, "", []):
                rows.append((label, self.md.meta[key]))
        known = {k for k, _ in S.COVER_KEYS} | set(S.PIPELINE_KEYS) | S.HIDDEN_KEYS
        for k, v in self.md.meta.items():
            if k not in known and v not in (None, "", []):
                rows.append((k, v))
        if rows:
            self._kv_table(self.doc, rows, S.BODY_W, S.COVER_LABEL_COL, S.LABEL_BG, S.TABLE_SIZE)

    def _kv_table(self, container, rows, width, label_ratio, label_bg, size, color=None):
        w0 = round(width * label_ratio)
        widths = [w0, width - w0]
        t = container.add_table(rows=len(rows), cols=2)
        self._table_props(t, widths)
        ctx = Ctx(width=widths[1], where="cell", pstyle="MDTable")
        for r, (label, value) in enumerate(rows):
            tr = t.rows[r]
            self._row_props(tr._tr, cant_split=True)
            c0, c1 = tr.cells
            self._cell_props(c0, widths[0], fill=label_bg)
            self._cell_props(c1, widths[1])
            p = self._new_par(c0, "MDTable")
            r_ = self._add_text(p._p, label, ctx=None)
            for rr in r_:
                ox.run_bold(rr)
                ox.run_size(rr, size)
                if color:
                    ox.run_color(rr, color)
            values = value if isinstance(value, list) else [value]
            for n, v in enumerate(values):
                p = self._new_par(c1, "MDTable")
                if len(values) > 1:
                    ox.numbering(p._p, self.bullet_num, 0)
                    ox.indent(p._p, left=S.CELL_LIST_INDENT, hanging=S.CELL_LIST_HANGING)
                for rr in self._add_text(p._p, self._val_text(v), ctx=None, emoji=True):
                    ox.run_size(rr, size)
                    if color:
                        ox.run_color(rr, color)
                if n < len(values) - 1:
                    ox.spacing(p._p, after=1)
            self._end_cell(c0)
            self._end_cell(c1)
        return t

    @staticmethod
    def _val_text(v):
        if isinstance(v, dict):
            return ", ".join(f"{k}: {x}" for k, x in v.items())
        return str(v)

    # ━━ 차례 ━━
    def _toc(self):
        p = self.doc.add_paragraph()
        ox.style(p._p, "MDTocTitle")
        ox.flag(p._p, "w:pageBreakBefore")
        self._add_text(p._p, S.TOC_TITLE, ctx=None)
        entries = [h for h in self.result.headings if h.level <= 2]
        avail = S.BODY_H / 20 - (S.TOC_TITLE_SIZE * 1.3 + 18 + 8) \
            - S.TOC_NOTE_SIZE * S.BODY_LINE - 36   # 안전 여유

        def heights(ents):
            hs = []
            for e in ents:
                if e.level == 1:
                    w, size, lh = S.BODY_W - 700, S.TOC_CH_SIZE, S.TOC_CH_SIZE * S.BODY_LINE
                else:
                    w, size, lh = S.BODY_W - 700 - S.LIST_INDENT, S.TOC_SUB_SIZE, S.TOC_SUB_SIZE * S.BODY_LINE
                hs.append(est_lines(e.text, w, size) * lh)
            return hs

        hs = heights(entries)
        if sum(hs) + 2 * len(entries) > avail:
            entries = [h for h in entries if h.level == 1]   # 소제목까지는 한 쪽에 안 들어간다
            hs = heights(entries)
        n_ch = sum(1 for e in entries if e.level == 1)
        n_sub = len(entries) - n_ch
        extra = max(0.0, avail - sum(hs))
        # 장 사이 간격은 소제목 간격의 2.5배
        units = 2.5 * max(n_ch - 1, 0) + n_sub
        g = extra / units if units else 0
        g = max(1.5, min(g, 9.0))
        used = 0.0
        for i, (e, hgt) in enumerate(zip(entries, hs)):
            p = self.doc.add_paragraph()
            ox.style(p._p, "MDToc1" if e.level == 1 else "MDToc2")
            before = 0 if i == 0 else (g * 2.5 if e.level == 1 else g)
            before = round(before * 2) / 2
            ox.spacing(p._p, before=before)
            used += before + hgt
            if e.level == 2:
                ox.indent(p._p, left=S.LIST_INDENT, right=0)
            ox.tabs(p._p, [("right", S.BODY_W, "dot")])
            hl = ox.internal_hyperlink(e.bookmark)
            p._p.append(hl)
            self._add_text(hl, e.text, ctx=None, emoji=True)
            hl.append(self._mk_run_el(tab=True))
            for r in ox.field_runs(f"PAGEREF {e.bookmark} \\h", ""):
                hl.append(r)
        p = self.doc.add_paragraph()
        ox.style(p._p, "MDTocNote")
        ox.spacing(p._p, before=max(12.0, round(avail - used)))
        self._add_text(p._p, S.TOC_NOTE, ctx=None)

    def _page_start(self):
        p = self.doc.add_paragraph()
        ox.style(p._p, "MDSpacer")
        ox.flag(p._p, "w:pageBreakBefore")
        ox.spacing(p._p, after=0, line=1)

    # ━━ 머리글·바닥글 ━━
    def _header_footer(self, title):
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        sec.first_page_header.is_linked_to_previous = False
        sec.first_page_footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        ox.style(hp._p, "Header")
        ox.align(hp._p, "right")
        self._add_text(hp._p, title, ctx=None, emoji=True)
        fp = sec.footer.paragraphs[0]
        ox.style(fp._p, "Footer")
        ox.align(fp._p, "center")
        for r in ox.field_runs("PAGE", "1"):
            fp._p.append(r)

    # ━━ 블록 렌더링 ━━
    def _render_blocks(self, container, blocks, ctx: Ctx):
        for i, b in enumerate(blocks):
            nxt = blocks[i + 1] if i + 1 < len(blocks) else None
            self._render_block(container, b, ctx, nxt)

    def _render_block(self, container, b, ctx: Ctx, nxt):
        if isinstance(b, Heading):
            if b is not self.title_heading:
                self._heading(container, b, ctx)
        elif isinstance(b, Paragraph):
            self._paragraph(container, b.spans, ctx)
        elif isinstance(b, ListBlock):
            self._list(container, b, ctx, 0)
            if self.last_p is not None:
                after = {"body": S.LIST_END_AFTER, "box": S.BOX_PARA_AFTER}.get(ctx.where, S.CELL_PARA_AFTER)
                ox.spacing(self.last_p, after=after)
        elif isinstance(b, Table):
            if b.header or b.rows:
                self._table(container, b, ctx)
                self._after_table(container, ctx, nxt)
        elif isinstance(b, Quote):
            self._box(container, b, ctx)
            self._after_table(container, ctx, nxt)
        elif isinstance(b, MathBlock):
            self._math_block(container, b, ctx)
        elif isinstance(b, CodeBlock):
            self._code(container, b, ctx)
        elif isinstance(b, Rule):
            if nxt is not None and not isinstance(b, Heading) and not isinstance(nxt, (Heading, Rule)) \
                    and ctx.where == "body":
                p = self._new_par(container, "MDSpacer")
                ox.spacing(p._p, after=S.HR_AFTER)

    def _after_table(self, container, ctx, nxt):
        if isinstance(nxt, Heading):
            return
        if nxt is None and ctx.where != "body":
            return
        self._new_par(container, "MDSpacer")

    # ── 단락 생성 ──
    def _new_par(self, container, pstyle):
        if isinstance(container, _Cell) and container._tc not in self.used_cells:
            self.used_cells.add(container._tc)
            p = container.paragraphs[0]
        else:
            p = container.add_paragraph()
        if pstyle != "Normal":
            ox.style(p._p, pstyle)
        self.last_p = p._p
        return p

    def _is_first_in(self, container):
        return isinstance(container, _Cell) and container._tc not in self.used_cells

    def _heading(self, container, h: Heading, ctx):
        lvl = self.levels.get(id(h), 3)
        p = self._new_par(container, f"Heading{lvl}")
        if ctx.base_left:
            ox.indent(p._p, left=ctx.base_left)
        self._add_spans(p._p, h.spans, ctx, heading=True)
        bm = self.bookmarks.get(id(h))
        if bm:
            ox.wrap_bookmark(p._p, bm, int(bm[4:]))

    def _is_label(self, spans) -> bool:
        txt = spans_text(spans).strip()
        if not txt or "\n" in txt or len(txt) > 50:
            return False
        texty = [s for s in spans if s.kind == "text" and s.text.strip()]
        if not texty or any(s.kind not in ("text",) for s in spans):
            return False
        if all(s.bold for s in texty):
            return True
        return len(txt) <= 20 and txt.endswith((":", "："))

    def _paragraph(self, container, spans, ctx: Ctx, indent_left=None, list_cont=False):
        """단락 하나(▶ 해설은 별도 단락으로 분리). 만든 단락 목록을 돌려준다."""
        made = []
        for kind, seg in split_feedback(spans):
            if not seg and kind == "normal":
                continue
            first = self._is_first_in(container)
            p = self._new_par(container, ctx.pstyle)
            made.append(p._p)
            left = indent_left if indent_left is not None else (ctx.base_left or None)
            if left:
                ox.indent(p._p, left=left)
            if kind == "feedback":
                rs = self._add_text(p._p, S.FEEDBACK_LABEL, ctx=None, emoji=True)
                for r in rs:
                    ox.run_bold(r)
                    ox.run_color(r, S.BOLD_COLOR)
                self._add_spans(p._p, seg, ctx)
                continue
            if not list_cont and self._is_label(seg):
                if ctx.where == "body":
                    ox.spacing(p._p, before=S.BODY_LABEL_BEFORE, after=S.BODY_LABEL_AFTER)
                else:
                    ox.spacing(p._p, before=0 if first else S.LABEL_BEFORE, after=S.LABEL_AFTER)
                ox.flag(p._p, "w:keepNext")
                self.label_ps.add(p._p)
            self._add_spans(p._p, seg, ctx)
            if any(s.kind in ("image", "math") for s in seg):
                ox.spacing(p._p, line=ox_line(ctx), rule="atLeast")
            if all(s.kind == "image" or (s.kind == "text" and not s.text.strip()) for s in seg):
                ox.align(p._p, "center")
        return made

    # ── 목록 ──
    def _list(self, container, lb: ListBlock, ctx: Ctx, level: int):
        in_cell = ctx.where == "cell"
        step = S.CELL_LIST_INDENT if in_cell else S.LIST_INDENT
        hang = S.CELL_LIST_HANGING if in_cell else S.LIST_HANGING
        ilvl = min(level, 8)
        if lb.ordered:
            num_id = self._add_num(self.abs_decimal, restart_level=ilvl, start=lb.start)
        else:
            num_id = self.bullet_num
        if lb.ordered and lb.start + len(lb.items) - 1 >= 10:
            hang += 120
            step += 120
        text_left = ctx.base_left + step * (level + 1)
        item_after = S.LIST_AFTER if lb.tight else S.LIST_LOOSE_AFTER
        for item in lb.items:
            first = True
            for blk in item.blocks:
                if first and isinstance(blk, Paragraph):
                    if item.checked is not None:
                        self._task_par(container, blk.spans, item.checked, ctx, level, step)
                        ps = [self.last_p]
                        # ▶ 해설 등 나머지는 일반 처리
                    else:
                        ps = self._paragraph(container, blk.spans, ctx, indent_left=text_left,
                                             list_cont=True)
                        if ps:
                            ox.numbering(ps[0], num_id, ilvl)
                            ox.indent(ps[0], left=text_left, hanging=hang)
                    for p in ps:
                        ox.spacing(p, after=item_after)
                elif isinstance(blk, ListBlock):
                    self._list(container, blk, ctx, level + 1)
                elif isinstance(blk, Paragraph):
                    for p in self._paragraph(container, blk.spans, ctx, indent_left=text_left,
                                             list_cont=True):
                        ox.spacing(p, after=item_after)
                else:
                    self._render_block(container, blk, ctx.sub(base_left=text_left,
                                                               width=ctx.width), None)
                first = False
            if first:   # 빈 항목
                p = self._new_par(container, ctx.pstyle)
                ox.numbering(p._p, num_id, ilvl)
                ox.indent(p._p, left=text_left, hanging=hang)
                ox.spacing(p._p, after=item_after)

    def _task_par(self, container, spans, checked, ctx, level, step):
        p = self._new_par(container, ctx.pstyle)
        left = ctx.base_left + step * level + S.CHECKBOX_HANGING
        ox.indent(p._p, left=left, hanging=S.CHECKBOX_HANGING)
        ox.tabs(p._p, [("left", left, None)])
        p._p.append(ox.checkbox_sdt(checked, self.sdt_ids.next(), S.SYMBOL_FONT, ctx.size))
        p._p.append(self._mk_run_el(tab=True))
        self._add_spans(p._p, [s for s in spans], ctx)
        return p

    # ── 표 ──
    def _col_widths(self, tbl: Table, total: int):
        """열 너비(twip). 내용 길이의 제곱근에 비례해 나누고, 가장 긴 단어가 꺾이지 않게 최소 폭을 둔다."""
        n = len(tbl.header) if tbl.header else len(tbl.rows[0])
        size = S.TABLE_SIZE
        pad = 2 * S.CELL_MAR_LR + 60
        scores, mins = [], []
        for c in range(n):
            hdr_txt = spans_text(tbl.header[c]) if tbl.header else ""
            cells = [spans_text(r[c]) for r in tbl.rows] or [""]
            lens = [max((text_em(x) for x in t.split("\n")), default=0) for t in cells]
            avg = sum(lens) / len(lens)
            mx = max(lens)
            body = 0.6 * avg + 0.4 * min(mx, 2 * avg + 5)
            scores.append(math.sqrt(max(body, text_em(hdr_txt) * 0.6, 1.0)))
            words = re.split(r"\s+", hdr_txt) + [w for t in cells for w in re.split(r"\s+", t)]
            longest = max((text_em(w) for w in words), default=1)
            mins.append(min(longest, 6.5) * size * 20 + pad)
        kind0 = self._first_col_kind(tbl)
        widths = [0.0] * n
        fixed0 = n >= 2 and kind0 in ("label", "narrow")
        if fixed0:
            if kind0 == "label":
                share = scores[0] / sum(scores)
                widths[0] = total * max(S.LABEL_COL_MIN, min(S.LABEL_COL_MAX, share))
            else:
                full = max(text_em(spans_text(r[0])) for r in tbl.rows)
                widths[0] = max(mins[0], full * size * 20 + pad + 60, total * 0.045)
            rest = sum(scores[1:])
            for c in range(1, n):
                widths[c] = (total - widths[0]) * scores[c] / rest
        else:
            tot = sum(scores)
            widths = [total * sc / tot for sc in scores]
        # 최소 폭 보정 (첫 고정 열 제외)
        start = 1 if fixed0 else 0
        for _ in range(4):
            deficit = 0.0
            for c in range(start, n):
                if widths[c] < mins[c]:
                    deficit += mins[c] - widths[c]
                    widths[c] = mins[c]
            if deficit <= 0.5:
                break
            big = [c for c in range(start, n) if widths[c] > mins[c] + 1]
            pool = sum(widths[c] - mins[c] for c in big)
            if pool <= 0:
                break
            for c in big:
                widths[c] -= deficit * (widths[c] - mins[c]) / pool
        w = [max(300, int(round(x))) for x in widths]
        w[-1] += total - sum(w)
        return w, kind0

    @staticmethod
    def _first_col_kind(tbl: Table):
        col = [text_em(spans_text(r[0])) for r in tbl.rows if r]
        if not col:
            return "text"
        mx, avg = max(col), sum(col) / len(col)
        if mx <= 2.5:
            return "narrow"
        if mx <= 24 and avg <= 16:
            return "label"
        return "text"

    def _table_props(self, t, widths, borders=True, indent=0):
        tbl = t._tbl
        tblPr = tbl.tblPr
        for c in list(tblPr):
            tblPr.remove(c)
        total = sum(widths)
        ox.set_child(tblPr, ox.el("w:tblW", w=total, type="dxa"))
        ox.set_child(tblPr, ox.el("w:tblInd", w=indent, type="dxa"))
        ox.tbl_borders(tblPr, sz=S.TABLE_BORDER_SZ, color=S.BORDER, none=not borders)
        ox.set_child(tblPr, ox.el("w:tblLayout", type="fixed"))
        ox.cell_margins(tblPr, S.CELL_MAR_TB, S.CELL_MAR_LR, S.CELL_MAR_TB, S.CELL_MAR_LR)
        ox.set_child(tblPr, ox.el("w:tblLook", val="04A0", firstRow=1, lastRow=0, firstColumn=1,
                                  lastColumn=0, noHBand=0, noVBand=1))
        grid = tbl.tblGrid
        for gc in list(grid):
            grid.remove(gc)
        for wv in widths:
            grid.append(ox.el("w:gridCol", w=wv))

    @staticmethod
    def _row_props(tr, cant_split=True, header=False):
        trPr = tr.get_or_add_trPr()
        for c in list(trPr):
            trPr.remove(c)
        if cant_split:
            trPr.append(ox.el("w:cantSplit"))
        if header:
            trPr.append(ox.el("w:tblHeader"))

    @staticmethod
    def _cell_props(cell, width, fill=None, valign=None):
        tcPr = cell._tc.get_or_add_tcPr()
        ox.set_child(tcPr, ox.el("w:tcW", w=width, type="dxa"))
        if fill:
            ox.shade(tcPr, fill)
        if valign:
            ox.set_child(tcPr, ox.el("w:vAlign", val=valign))

    def _table(self, container, tbl: Table, ctx: Ctx):
        width = ctx.width - ctx.base_left
        widths, kind0 = self._col_widths(tbl, width)
        n = len(widths)
        nrows = (1 if tbl.header else 0) + len(tbl.rows)
        t = self._add_table(container, nrows, n)
        self._table_props(t, widths, indent=ctx.base_left)
        r0 = 0
        rows_el = t.rows
        if tbl.header and any(spans_text(c).strip() for c in tbl.header):
            tr = rows_el[0]
            self._row_props(tr._tr, cant_split=True, header=True)
            hdr_bg = S.BOX_KINDS[tbl.accent][1] if tbl.accent in S.BOX_KINDS else S.HEADER_ROW_BG
            for c, cell in enumerate(tr.cells):
                self._cell_props(cell, widths[c], fill=hdr_bg, valign="center")
                p = self._new_par(cell, "MDTable")
                ox.align(p._p, "center")
                ox.spacing(p._p, after=0)
                self._add_spans(p._p, tbl.header[c], ctx.sub(where="cell", width=widths[c]),
                                header_row=True)
            r0 = 1
        elif tbl.header:
            t._tbl.remove(rows_el[0]._tr)
            rows_el = t.rows
        page_third = S.BODY_H / 20 * S.ROW_SPLIT_RATIO
        for ri, row in enumerate(tbl.rows):
            tr = rows_el[r0 + ri]
            cells = tr.cells
            row_h = 0.0
            for c, cell in enumerate(cells):
                label = c == 0 and kind0 == "label" and n >= 2
                self._cell_props(cell, widths[c], fill=S.FIRST_COL_BG if label else None)
                cctx = Ctx(width=widths[c] - 2 * S.CELL_MAR_LR, where="cell", pstyle="MDTable")
                self._cell_content(cell, row[c], cctx, label=label,
                                   align=tbl.aligns[c] if c < len(tbl.aligns) else None,
                                   center=(c == 0 and kind0 == "narrow"))
                row_h = max(row_h, self._est_height(cell._tc, cctx.width))
            self._row_props(tr._tr, cant_split=row_h <= page_third)
        return t

    def _add_table(self, container, rows, cols):
        """표를 추가한다. 셀 안이면 python-docx가 붙이는 빈 단락과 비어 있던 첫 단락을 없앤다."""
        if not isinstance(container, _Cell):
            return container.add_table(rows=rows, cols=cols)
        was_first = self._is_first_in(container)
        t = container.add_table(rows=rows, cols=cols)
        tbl = t._tbl
        nxt = tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p") and not len(nxt):
            container._tc.remove(nxt)
        if was_first:
            prev = tbl.getprevious()
            if prev is not None and prev.tag == qn("w:p") and prev.find(qn("w:r")) is None:
                container._tc.remove(prev)
            self.used_cells.add(container._tc)
        return t

    def _cell_content(self, cell, spans, ctx: Ctx, label=False, align=None, center=False):
        groups = cell_groups(spans)
        feedback = [g for k, g in groups if k == "feedback"]
        normal = [(k, g) for k, g in groups if k != "feedback"]
        made = []
        for k, g in normal:
            p = self._new_par(cell, "MDTable")
            made.append(p._p)
            if k == "bullet":
                ox.numbering(p._p, self.bullet_num, 0)
                ox.indent(p._p, left=S.CELL_LIST_INDENT, hanging=S.CELL_LIST_HANGING)
                ox.spacing(p._p, after=1)
            self._add_spans(p._p, g, ctx, label_cell=label)
            if any(s.kind in ("image", "math") for s in g):
                ox.spacing(p._p, line=ox_line(ctx), rule="atLeast")
            jc = {"center": "center", "right": "right"}.get(align or "")
            if center:
                jc = "center"
            if jc and k != "bullet":
                ox.align(p._p, jc)
        for g in feedback:
            p = self._new_par(cell, "MDTable")
            made.append(p._p)
            ox.spacing(p._p, before=4 if made[:-1] else 0)
            for r in self._add_text(p._p, S.FEEDBACK_LABEL, ctx=None, emoji=True):
                ox.run_bold(r)
                ox.run_color(r, S.BOLD_COLOR)
            self._add_spans(p._p, g, ctx)
        if not made:
            self._new_par(cell, "MDTable")
        self._end_cell(cell)

    def _end_cell(self, cell):
        """셀은 단락으로 끝나야 하고, 마지막 단락 뒤 간격은 0."""
        tc = cell._tc
        last = tc[-1]
        if last.tag == qn("w:tbl"):
            p = cell.add_paragraph()
            ox.style(p._p, "MDSpacer")
            ox.spacing(p._p, after=0, line=1)
            return
        if last.tag == qn("w:p"):
            ox.spacing(last, after=0)

    def _est_height(self, tc, width) -> float:
        h = 0.0
        for p in tc.iter(qn("w:p")):
            sid = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
            sid = sid.get(qn("w:val")) if sid is not None else "Normal"
            if sid == "MDTable":
                size, lh = S.TABLE_SIZE, S.TABLE_SIZE * S.TABLE_LINE
            elif sid == "MDSpacer":
                h += S.SPACER_LINE + S.SPACER_AFTER
                continue
            else:
                size, lh = S.BODY_SIZE, S.BODY_SIZE * S.BODY_LINE
            txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
            nbr = sum(1 for _ in p.iter(qn("w:br")))
            h += (est_lines(txt, width, size) + nbr) * lh + (ox.get_after(p) or 0)
        return h + 2 * S.CELL_MAR_TB / 20

    # ── 콜아웃 상자 ──
    def _box(self, container, q: Quote, ctx: Ctx):
        prefix, bar, bg = S.BOX_KINDS.get(q.kind, S.BOX_KINDS["quote"])
        width = ctx.width - ctx.base_left
        t = self._add_table(container, 1, 1)
        self._table_props(t, [width], borders=False, indent=ctx.base_left)
        tr = t.rows[0]
        cell = tr.cells[0]
        self._cell_props(cell, width, fill=bg)
        tcPr = cell._tc.get_or_add_tcPr()
        ox.tc_borders(tcPr, left=("single", S.BOX_BAR_SZ, bar))
        m = S.BOX_MAR
        ox.cell_margins(tcPr, m["top"], m["left"], m["bottom"], m["right"], tag="w:tcMar")
        inner = Ctx(width=width - m["left"] - m["right"] - 60, where="box", pstyle="MDBox")
        title_p = None
        blocks = q.blocks
        # 괄호 없는 아이콘 접두어(📌, ✅)만 있고 제목이 없으면 첫 문장 앞에 붙인다
        if prefix and not q.title and not prefix.startswith("[") and blocks                 and isinstance(blocks[0], Paragraph) and not self._is_label(blocks[0].spans):
            blocks = [Paragraph([Span(prefix + " ")] + blocks[0].spans)] + blocks[1:]
            prefix = ""
        if prefix or q.title:
            title_p = self._new_par(cell, "MDBox")
            ox.spacing(title_p._p, after=4)
            rs = []
            if prefix:
                rs += self._add_text(title_p._p, prefix + (" " if q.title else ""), ctx=None, emoji=True)
            if q.title:
                rs += self._add_spans(title_p._p, q.title, inner, plain=True)
            for r in rs:
                ox.run_bold(r)
                ox.run_color(r, bar if q.kind != "quote" else S.DARK_TEXT)
        self._render_blocks(cell, blocks, inner)
        if not self.used_cells.__contains__(cell._tc):
            self._new_par(cell, "MDBox")
        self._end_cell(cell)
        self._box_rows(t, cell, inner.width, m)

    def _box_rows(self, t, cell, inner_w, m):
        """상자의 쪽 나눔 규칙을 행 구조로 구현한다.

        표 셀 안의 keepNext는 Word에서 「행 전체를 다음과 묶음」으로 동작해 상자가 통째로
        다음 쪽으로 밀린다. 그래서 keepNext 대신 행을 나눈다.
        - 6줄 이하: 분할 금지 행 하나
        - 길면: 제목+첫 두 항목(분할 금지 행) / 나머지(분할 허용 행).
          상자 안 소제목성 레이블마다 같은 방식으로 새 행을 시작한다.
        행 사이에는 테두리가 없고 막대·배경이 이어져 하나의 상자로 보인다.
        """
        tc = cell._tc
        kids = [c for c in tc if c.tag in (qn("w:p"), qn("w:tbl"))]
        for p in tc.iter(qn("w:p")):
            pPr = p.find(qn("w:pPr"))
            if pPr is not None:
                for tag in ("w:keepNext", "w:keepLines"):
                    e = pPr.find(qn(tag))
                    if e is not None:
                        pPr.remove(e)

        def lines(elems):
            n = 0
            for e in elems:
                if e.tag == qn("w:tbl"):
                    n += 99
                    continue
                txt = "".join(x.text or "" for x in e.iter(qn("w:t")))
                n += est_lines(txt, inner_w, S.BODY_SIZE) + sum(1 for _ in e.iter(qn("w:br")))
            return n

        if lines(kids) <= S.BOX_KEEP_LINES:
            self._row_props(t.rows[0]._tr, cant_split=True)
            return
        # 레이블 단락에서 묶음을 나눈다
        chunks, cur = [], []
        for e in kids:
            if cur and e in self.label_ps:
                chunks.append(cur)
                cur = []
            cur.append(e)
        chunks.append(cur)
        specs = []
        for ch in chunks:
            if lines(ch) <= S.BOX_KEEP_LINES + 2 or len(ch) <= 3:
                specs.append((ch, True))
                continue
            k = 3
            while k > 1 and ch[k - 1].tag != qn("w:p"):
                k -= 1
            specs.append((ch[:k], True))
            specs.append((ch[k:], False))
        specs = [(ch, cs) for ch, cs in specs if ch]
        tcPr0 = tc.find(qn("w:tcPr"))
        rows = [t.rows[0]._tr]
        for _ in specs[1:]:
            tr = deepcopy_el(rows[0])
            new_tc = tr.find(qn("w:tc"))
            for c in list(new_tc):
                if c.tag != qn("w:tcPr"):
                    new_tc.remove(c)
            t._tbl.append(tr)
            rows.append(tr)
        last = len(specs) - 1
        for i, ((elems, cant), tr) in enumerate(zip(specs, rows)):
            row_tc = tr.find(qn("w:tc"))
            if i:
                for e in elems:
                    row_tc.append(e)
            if row_tc[-1].tag != qn("w:p"):
                sp = parse_xml('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
                ox.style(sp, "MDSpacer")
                ox.spacing(sp, after=0, line=1)
                row_tc.append(sp)
            tcPr = row_tc.find(qn("w:tcPr")) if i else tcPr0
            ox.cell_margins(tcPr, m["top"] if i == 0 else 0, m["left"],
                            m["bottom"] if i == last else 0, m["right"], tag="w:tcMar")
            self._row_props(tr, cant_split=cant)

    # ── 수식 ──
    def _math_block(self, container, b: MathBlock, ctx):
        p = self._new_par(container, ctx.pstyle)
        ox.align(p._p, "center")
        ox.spacing(p._p, before=4, after=8, line=ox_line(ctx), rule="atLeast")
        if ctx.base_left:
            ox.indent(p._p, left=ctx.base_left)
        omml = self._omml(b.latex)
        if omml is not None:
            para = parse_xml('<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"/>')
            para.append(omml)
            p._p.append(para)
        else:
            self._math_fallback(p._p, b.latex)

    def _omml(self, latex):
        if not self.opts.mml2omml or self._xslt_failed:
            return None
        try:
            from lxml import etree
            import latex2mathml.converter
            if self._xslt is None:
                self._xslt = etree.XSLT(etree.parse(self.opts.mml2omml))
            mathml = latex2mathml.converter.convert(latex)
            res = self._xslt(etree.fromstring(mathml.encode("utf-8")))
            root = res.getroot()
            if root is None:
                return None
            M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
            if root.tag == f"{{{M}}}oMathPara":
                inner = root.find(f"{{{M}}}oMath")
                root = inner if inner is not None else root
            omml = parse_xml(etree.tostring(root))
            # 한글이 들어간 수식 글자는 맑은 고딕으로
            for mr in omml.iter(f"{{{M}}}r"):
                rpr = ox.el("w:rPr")
                rpr.append(ox.el("w:rFonts", ascii=S.MATH_FONT, hAnsi=S.MATH_FONT,
                                 eastAsia=S.FONT, cs=S.MATH_FONT))
                mrpr = mr.find(f"{{{M}}}rPr")
                if mrpr is not None:
                    mrpr.addnext(rpr)
                else:
                    mr.insert(0, rpr)
            return omml
        except OSError:
            self._xslt_failed = True
            return None
        except Exception as e:  # noqa: BLE001 - 수식 하나 실패는 대체 표시로
            self.result.warnings.append(f"수식 변환 실패, 텍스트로 대체: {latex[:40]} ({e})")
            return None

    def _math_fallback(self, parent, latex):
        txt = latex_plain(latex)
        r = self._mk_run_el(txt)
        ox.run_fonts(r, S.MATH_FONT)
        ox.set_child(r.find(qn("w:rPr")), ox.el("w:rFonts", ascii=S.MATH_FONT, hAnsi=S.MATH_FONT,
                                                eastAsia=S.FONT, cs=S.MATH_FONT))
        ox.run_italic(r)
        parent.append(r)
        return r

    # ── 코드 ──
    def _code(self, container, b: CodeBlock, ctx):
        p = self._new_par(container, "MDCode")
        if ctx.base_left:
            ox.indent(p._p, left=ctx.base_left)
        ox.spacing(p._p, before=2, after=8 if ctx.where == "body" else 4)
        lines = b.text.replace("\t", "    ").split("\n")
        for i, line in enumerate(lines):
            if i:
                br = self._mk_run_el()
                br.append(ox.el("w:br"))
                p._p.append(br)
            p._p.append(self._mk_run_el(line))

    # ── 각주·부록 ──
    def _footnotes(self, ctx):
        if not self.md.footnotes:
            return
        p = self._new_par(self.doc, "Normal")
        ox.spacing(p._p, before=S.BODY_LABEL_BEFORE, after=S.BODY_LABEL_AFTER)
        ox.flag(p._p, "w:keepNext")
        for r in self._add_text(p._p, S.FOOTNOTE_TITLE, ctx=None):
            ox.run_bold(r)
            ox.run_color(r, S.BOLD_COLOR)
        for fn in self.md.footnotes:
            spans = [Span(f"[{fn.label}] ")]
            for blk in fn.blocks:
                if isinstance(blk, Paragraph):
                    spans += blk.spans + [Span(" ")]
            q = self._new_par(self.doc, "Normal")
            ox.spacing(q._p, after=2)
            for r in self._add_spans(q._p, spans, ctx):
                ox.run_size(r, S.FOOTNOTE_SIZE)

    def _appendix_items(self):
        return [(label, self.md.meta[k]) for k, label in S.PIPELINE_KEYS.items()
                if self.md.meta.get(k) not in (None, "", [])]

    def _appendix(self, ctx):
        items = self._appendix_items()
        if not items:
            return
        p = self._new_par(self.doc, "Heading1")
        self._add_text(p._p, S.APPENDIX_TITLE, ctx=None)
        ox.wrap_bookmark(p._p, self._appendix_bm, int(self._appendix_bm[4:]))
        self._kv_table(self.doc, items, S.BODY_W, 0.22, S.LABEL_BG, S.APPENDIX_SIZE,
                       color=S.GRAY_TEXT)

    def _trim_trailing(self):
        body = self.doc.element.body
        sect = body.find(qn("w:sectPr"))
        prev = sect.getprevious() if sect is not None else None
        # python-docx 기본 문서의 빈 첫 단락이 있으면 제거
        first = body[0]
        if first.tag == qn("w:p") and not len(first.findall(".//" + qn("w:r"))) \
                and first.find(qn("w:pPr") + "/" + qn("w:pStyle")) is None:
            body.remove(first)
        return prev

    # ━━ 인라인 ━━
    def _mk_run_el(self, text=None, tab=False):
        r = parse_xml('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        if tab:
            r.append(ox.el("w:tab"))
        if text is not None:
            t = ox.el("w:t")
            t.text = text
            t.set(qn("xml:space"), "preserve")
            r.append(t)
        return r

    def _split_emoji(self, text):
        segs = split_emoji(text)
        if self.opts.for_pdf:
            # U+FE0F(이모지 표시 선택자)가 있으면 Word가 Segoe UI Emoji로 대체해 통째로 넣는다
            segs = [(t.replace("️", ""), e) if e else (t, e) for t, e in segs]
            segs = [(t, e) for t, e in segs if t] or [("", False)]
        return segs

    def _add_text(self, parent, text, ctx, emoji=True):
        """서식 없는 텍스트를 이모지 글꼴 분리만 해서 넣는다."""
        runs = []
        for seg, is_emoji in self._split_emoji(text) if emoji else [(text, False)]:
            r = self._mk_run_el(seg)
            if is_emoji:
                ox.run_fonts(r, self.emoji_font)
            parent.append(r)
            runs.append(r)
        return runs

    def _add_spans(self, parent, spans, ctx: Ctx | None, heading=False, header_row=False,
                   label_cell=False, plain=False):
        """Span 목록을 런으로 바꿔 parent(w:p 또는 w:hyperlink)에 붙인다."""
        runs = []
        in_table = ctx is not None and ctx.where == "cell"
        hl_cache = {}
        for s in spans:
            if s.kind == "br":
                r = self._mk_run_el()
                r.append(ox.el("w:br"))
                parent.append(r)
                continue
            if s.kind == "image":
                self._image(parent, s, ctx)
                continue
            if s.kind == "math":
                om = self._omml(s.src)
                if om is not None:
                    parent.append(om)
                else:
                    runs.append(self._math_fallback(parent, s.src))
                continue
            if s.kind == "checkbox":
                continue
            target = parent
            if s.href and not s.href.startswith("#") and not heading:
                hl = hl_cache.get(s.href)
                if hl is None or hl.getparent() is None or parent[-1] is not hl:
                    hl = ox.external_hyperlink(self.part, s.href)
                    parent.append(hl)
                    hl_cache[s.href] = hl
                target = hl
            text = s.text
            if s.code:
                r = self._mk_run_el(text)
                ox.run_style(r, "MDCodeChar")
                target.append(r)
                runs.append(r)
                continue
            if not in_table:
                text = self._clean_text(text)
            pieces = self._decorate(text, allow_num=(self.opts.auto_number_bold and not heading
                                                     and not header_row and not s.bold
                                                     and not plain and not s.href))
            for piece, extra_bold, color in pieces:
                for seg, is_emoji in self._split_emoji(piece):
                    r = self._mk_run_el(seg)
                    bold = (s.bold or extra_bold) and not heading
                    if is_emoji:
                        ox.run_fonts(r, self.emoji_font)
                    if header_row:
                        ox.run_bold(r)
                        ox.run_color(r, S.HEADER_ROW_TEXT)
                    elif label_cell:
                        ox.run_bold(r)
                        ox.run_color(r, color or S.BOLD_COLOR)
                    elif bold and not plain:
                        ox.run_bold(r)
                        ox.run_color(r, color or S.BOLD_COLOR)
                    elif color:
                        ox.run_color(r, color)
                    if s.italic:
                        ox.run_italic(r)
                    if s.strike:
                        ox.run_flag(r, "w:strike")
                    if s.sup:
                        ox.set_child(ox.get_rPr(r), ox.el("w:vertAlign", val="superscript"))
                    if s.href and target is not parent:
                        ox.run_style(r, "MDLink")
                    target.append(r)
                    runs.append(r)
        return runs

    def _clean_text(self, text: str) -> str:
        """표 밖 텍스트의 마크다운 잔재 정리: ` — ` → `: `, 파이프 제거."""
        text = EMDASH_RE.sub(": ", text)
        if "|" in text:
            text = PIPE_RE.sub(" ", text)
        return text

    @staticmethod
    def _decorate(text, allow_num=True):
        """[?] 주황색, 숫자+단위 굵게. [(조각, 굵게, 색)] 목록."""
        marks = []
        for m in UNCERTAIN_RE.finditer(text):
            marks.append((m.start(), m.end(), False, S.UNCERTAIN_COLOR))
        if allow_num:
            for m in NUM_RE.finditer(text):
                frag = m.group(0)
                if re.fullmatch(r"\d{4}\s*년", frag):
                    continue
                if frag.endswith("일") and re.search(r"월\s*$", text[:m.start()]):
                    continue
                if any(a < m.end() and m.start() < b for a, b, _, _ in marks):
                    continue
                marks.append((m.start(), m.end(), True, None))
        if not marks:
            return [(text, False, None)]
        marks.sort()
        out, pos = [], 0
        for a, b, bold, color in marks:
            if a > pos:
                out.append((text[pos:a], False, None))
            out.append((text[a:b], bold, color))
            pos = b
        if pos < len(text):
            out.append((text[pos:], False, None))
        return out

    # ── 이미지 ──
    def _image(self, parent, s: Span, ctx):
        src = unquote(s.src or "")
        name = os.path.basename(src) or s.text or "?"
        path = None
        if src and not re.match(r"^[a-z]+://", src, re.I):
            cand = src if os.path.isabs(src) else os.path.join(self.md_dir, src)
            if os.path.exists(cand):
                path = cand
        if path is None:
            r = self._mk_run_el(f"[이미지 없음: {name}]")
            ox.run_color(r, S.GRAY_TEXT)
            parent.append(r)
            return
        try:
            from docx.image.image import Image as DocxImage
            from docx.shared import Emu
            img = DocxImage.from_file(path)
            native = img.px_width / (img.horz_dpi or 96) * 914400
            avail = ((ctx.width - ctx.base_left) if ctx else S.BODY_W) * 635
            width = Emu(int(min(native, avail)))
            from docx.text.paragraph import Paragraph as DocxPar
            p = parent if parent.tag == qn("w:p") else parent.getparent()
            run = DocxPar(p, None).add_run()
            run.add_picture(path, width=width)
            if parent is not p:
                parent.append(run._r)
        except Exception as e:  # noqa: BLE001
            r = self._mk_run_el(f"[이미지 삽입 실패: {name}]")
            ox.run_color(r, S.GRAY_TEXT)
            parent.append(r)
            self.result.warnings.append(f"이미지 삽입 실패: {name} ({e})")

    # ━━ 잔재 검사 ━━
    def _residue_check(self):
        found = {}
        parts = [self.doc.element]
        sec = self.doc.sections[0]
        parts += [sec.header._element, sec.footer._element]
        for root in parts:
            for t in root.iter(qn("w:t")):
                r = t.getparent()
                rs = r.find(qn("w:rPr") + "/" + qn("w:rStyle"))
                if rs is not None and rs.get(qn("w:val")) == "MDCodeChar":
                    continue
                p = r.getparent()
                while p is not None and p.tag != qn("w:p"):
                    p = p.getparent()
                if p is not None:
                    ps = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
                    if ps is not None and ps.get(qn("w:val")) == "MDCode":
                        continue
                txt = t.text or ""
                for pat in RESIDUE:
                    if pat in txt and pat not in found:
                        i = txt.index(pat)
                        found[pat] = txt[max(0, i - 15): i + 15]
        return [f"마크다운 잔재 '{k}' 발견: …{v}…" for k, v in found.items()]


# ── 모듈 함수 ──────────────────────────────────────────
def deepcopy_el(e):
    return copy.deepcopy(e)


def ox_line(ctx: Ctx | None):
    if ctx is not None and ctx.where == "cell":
        return S.TABLE_SIZE * S.TABLE_LINE
    return S.BODY_SIZE * S.BODY_LINE


def split_emoji(text: str):
    out, pos = [], 0
    for m in EMOJI_RE.finditer(text):
        seg = m.group(0)
        if all(ch in _ARROWS for ch in seg):
            continue
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((seg, True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    # 화살표만 이모지 범위에서 빼면서 생긴 인접 일반 조각 합치기
    merged = []
    for seg, e in out:
        if merged and not e and not merged[-1][1]:
            merged[-1] = (merged[-1][0] + seg, False)
        else:
            merged.append((seg, e))
    return merged or [("", False)]


def _split_spans_at(spans, pattern, keep_marker=False):
    """정규식 경계에서 Span 목록을 나눈다. 경계 문자열은 버린다."""
    groups, cur = [], []
    for s in spans:
        if s.kind != "text" or s.code:
            cur.append(s)
            continue
        pos = 0
        for m in pattern.finditer(s.text):
            if m.start() == 0 and m.end() == 0:
                continue
            head = s.text[pos:m.start()]
            if head:
                cur.append(replace(s, text=head))
            groups.append(cur)
            cur = []
            pos = m.end()
        tail = s.text[pos:]
        if tail:
            cur.append(replace(s, text=tail))
    groups.append(cur)
    return groups


def _lines(spans):
    lines, cur = [], []
    for s in spans:
        if s.kind == "br":
            lines.append(cur)
            cur = []
        else:
            cur.append(s)
    lines.append(cur)
    return lines


def _strip(spans):
    spans = list(spans)
    while spans and spans[0].kind == "text" and not spans[0].text.strip():
        spans.pop(0)
    while spans and spans[-1].kind == "text" and not spans[-1].text.strip():
        spans.pop()
    if spans and spans[0].kind == "text":
        spans[0] = replace(spans[0], text=spans[0].text.lstrip())
    if spans and spans[-1].kind == "text":
        spans[-1] = replace(spans[-1], text=spans[-1].text.rstrip())
    return spans


def _split_line_feedback(line):
    """한 줄을 ▶ 위치에서 나눈다: [(normal|feedback, spans)]"""
    out = []
    starts_fb = bool(line) and line[0].kind == "text" and line[0].text.lstrip().startswith("▶")
    parts = _split_spans_at(line, FEED_RE)
    for n, part in enumerate(parts):
        if n == 0 and not starts_fb:
            if _strip(part):
                out.append(("normal", _strip(part)))
        else:
            if _strip(part):
                out.append(("feedback", _strip(part)))
    return out


def split_feedback(spans):
    """본문 단락용: 일반 줄은 줄바꿈을 유지한 채 묶고, ▶ 해설은 따로 뗀다."""
    if not any(s.kind == "text" and "▶" in s.text for s in spans):
        return [("normal", spans)]
    out = []
    for line in _lines(spans):
        for kind, seg in _split_line_feedback(line):
            if kind == "normal" and out and out[-1][0] == "normal":
                out[-1] = ("normal", out[-1][1] + [Span(kind="br")] + seg)
            else:
                out.append((kind, seg))
    return out


def cell_groups(spans):
    """표 셀용: <br>은 단락, 긴 여러 문장은 글머리 기호, ▶는 해설. [(normal|bullet|feedback, spans)]"""
    out = []
    for line in _lines(spans):
        for kind, seg in _split_line_feedback(line):
            if kind == "feedback":
                out.append((kind, seg))
                continue
            text = spans_text(seg)
            if len(text) > S.LONG_CELL_CHARS:
                sents = [_strip(g) for g in _split_spans_at(seg, SENT_RE)]
                sents = [g for g in sents if g]
                merged = []
                for g in sents:
                    if merged and len(spans_text(g)) < 6:
                        merged[-1] = merged[-1] + [Span(" ")] + g
                    else:
                        merged.append(g)
                if len(merged) >= 2:
                    out.extend(("bullet", g) for g in merged)
                    continue
            out.append(("normal", seg))
    return out


_LATEX_SYMS = {r"\rightarrow": "→", r"\to": "→", r"\leftarrow": "←", r"\Rightarrow": "⇒",
               r"\times": "×", r"\cdot": "·", r"\le": "≤", r"\ge": "≥", r"\neq": "≠",
               r"\pm": "±", r"\approx": "≈", r"\infty": "∞", r"\quad": " ", r"\,": " "}


def latex_plain(latex: str) -> str:
    s = latex
    s = re.sub(r"\\underbrace\{((?:[^{}]|\{[^{}]*\})*)\}_\{((?:[^{}]|\{[^{}]*\})*)\}", r"\1(\2)", s)
    s = re.sub(r"\\(?:text|mathrm|mathbf|textbf|operatorname)\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    for k in sorted(_LATEX_SYMS, key=len, reverse=True):
        s = s.replace(k, _LATEX_SYMS[k])
    s = re.sub(r"\\([A-Za-z]+)", r"\1", s)
    s = s.replace("{", "").replace("}", "")
    return re.sub(r"\s+", " ", s).strip()


def build_docx(md: MdDocument, md_path: str, out_path: str, opts: BuildOptions | None = None):
    return DocxBuilder(md, md_path, opts).build(out_path)
